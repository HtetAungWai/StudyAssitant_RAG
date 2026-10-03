"""Turn uploaded files into text units, then into overlapping chunks.

Each unit keeps a 'location' (Page 3, Slide 7) so answers can cite sources.
"""
from io import BytesIO

SUPPORTED_TYPES = ["pdf", "pptx", "docx", "txt", "md"]


def parse_file(filename: str, data: bytes) -> list[dict]:
    """Return a list of {'text': str, 'location': str} units for a file."""
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        return _parse_pdf(data)
    if ext == "pptx":
        return _parse_pptx(data)
    if ext == "docx":
        return _parse_docx(data)
    if ext in ("txt", "md"):
        return [{"text": data.decode("utf-8", errors="ignore"), "location": "Text"}]
    raise ValueError(f"Unsupported file type: .{ext}")


def _parse_pdf(data: bytes) -> list[dict]:
    import pymupdf as fitz

    units = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        for i, page in enumerate(doc, start=1):
            text = page.get_text().strip()
            if text:
                units.append({"text": text, "location": f"Page {i}"})
    return units


def _parse_pptx(data: bytes) -> list[dict]:
    from pptx import Presentation

    units = []
    prs = Presentation(BytesIO(data))
    for i, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append("Speaker notes: " + notes)
        text = "\n".join(parts).strip()
        if text:
            units.append({"text": text, "location": f"Slide {i}"})
    return units


def _parse_docx(data: bytes) -> list[dict]:
    from docx import Document

    doc = Document(BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return [{"text": text, "location": "Document"}] if text else []


def chunk_text(text: str, size: int = 800, overlap: int = 150) -> list[str]:
    """Split text into ~size-character chunks with overlap, breaking on spaces."""
    text = " ".join(text.split())
    chunks, start = [], 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            space = text.rfind(" ", start + size // 2, end)
            if space != -1:
                end = space
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return chunks


def file_to_chunks(filename: str, data: bytes) -> list[dict]:
    """Parse + chunk a file. Returns [{'text', 'source', 'location'}, ...]."""
    out = []
    for unit in parse_file(filename, data):
        for piece in chunk_text(unit["text"]):
            out.append({"text": piece, "source": filename, "location": unit["location"]})
    return out
